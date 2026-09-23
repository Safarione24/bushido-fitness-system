from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from django.utils import timezone
from .models import Booking


def export_bookings_to_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Записи'

    headers = [
        'ID', 'Пользователь', 'Email', 'Сессия', 'Тренер',
        'Зона', 'Дата сессии', 'Дата записи', 'Статус', 'Причина'
    ]
    ws.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1E3A5F')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for b in Booking.objects.select_related('user', 'session', 'session__trainer', 'session__zone'):
        ws.append([
            b.pk,
            b.user.username,
            b.user.email or '—',
            b.session.title,
            str(b.session.trainer),
            b.session.zone.name,
            b.session.date.strftime('%d.%m.%Y %H:%M'),
            b.created_at.strftime('%d.%m.%Y %H:%M'),
            b.get_status_display(),
            b.moderation_reason or '—',
        ])

    widths = [6, 18, 24, 22, 22, 16, 18, 18, 14, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    ws.freeze_panes = 'A2'

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_bookings_by_session(session):
    wb = Workbook()
    ws = wb.active
    ws.title = f'Сессия {session.pk}'

    ws.append(['Сессия', session.title])
    ws.append(['Дата', session.date.strftime('%d.%m.%Y %H:%M')])
    ws.append(['Тренер', str(session.trainer)])
    ws.append(['Зона', session.zone.name])
    ws.append([])

    headers = ['№', 'Пользователь', 'Email', 'Дата записи', 'Статус', 'Причина']
    ws.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1E3A5F')
    for cell in ws[6]:
        cell.font = header_font
        cell.fill = header_fill

    for i, b in enumerate(session.booking_set.select_related('user'), 1):
        ws.append([
            i,
            b.user.username,
            b.user.email or '—',
            b.created_at.strftime('%d.%m.%Y %H:%M'),
            b.get_status_display(),
            b.moderation_reason or '—',
        ])

    for i, w in enumerate([6, 20, 26, 20, 14, 30], 1):
        ws.column_dimensions[ws.cell(row=6, column=i).column_letter].width = w

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer